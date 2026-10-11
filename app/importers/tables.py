from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

MAX_ROWS = 200_000


@dataclass(frozen=True)
class MapField:
    key: str
    label: str
    required: bool
    guesses: tuple[str, ...]


DNA_FIELDS = (
    MapField("name", "Имя", True, ("имя", "name", "match", "матч", "фио")),
    MapField("shared_cm", "Общий cM", False, ("shared cm", "общий cm", "общий см", "cm", "см")),
    MapField("largest_segment", "Крупнейший сегмент", False, ("segment", "сегмент", "largest", "крупн")),
)

CLUSTER_FIELDS = (
    MapField("cluster", "Номер кластера", True, ("cluster", "кластер")),
    MapField("name", "Имя", True, ("имя", "name", "match", "матч", "фио")),
    MapField("shared_cm", "Общий cM", False, ("shared cm", "общий cm", "общий см", "cm", "см")),
)

LINK_FIELDS = (
    MapField("name", "Имя", True, ("имя", "name", "фио")),
    MapField("birth_year", "Год рождения", False, ("год", "year", "birth", "рожд")),
)

FIELDS = {
    "matches": DNA_FIELDS,
    "clusters": CLUSTER_FIELDS,
    "links": LINK_FIELDS,
}


def parse_float(value: str) -> float | None:
    raw = value.strip().replace("\xa0", "").replace(" ", "")
    raw = re.sub(r"(?i)(cm|см)", "", raw)
    if not raw:
        return None
    if "," in raw and "." in raw:
        if raw.rfind(",") > raw.rfind("."):
            raw = raw.replace(".", "").replace(",", ".")
        else:
            raw = raw.replace(",", "")
    else:
        raw = raw.replace(",", ".")
    raw = re.sub(r"[^0-9.\-]", "", raw)
    if not raw or raw in {"-", ".", "-."}:
        return None
    try:
        return float(raw)
    except ValueError:
        return None


def cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def unique_headers(cells: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(cells, start=1):
        name = " ".join(cell.split()) or f"Столбец {index}"
        count = seen.get(name, 0) + 1
        seen[name] = count
        headers.append(name if count == 1 else f"{name} ({count})")
    return headers


def table_from_rows(raw_rows: list[list[str]]) -> tuple[list[str], list[dict[str, str]]]:
    if not raw_rows:
        raise ValueError("В таблице нет строк.")
    headers = unique_headers(raw_rows[0])
    records: list[dict[str, str]] = []
    for row in raw_rows[1:]:
        padded = row + [""] * (len(headers) - len(row))
        item = {headers[index]: padded[index].strip() for index in range(len(headers))}
        if any(item.values()):
            records.append(item)
    if not records:
        raise ValueError("В таблице нет строк.")
    if len(records) > MAX_ROWS:
        raise ValueError("В таблице больше 200 000 строк.")
    return headers, records


def detect_delimiter(sample: str) -> str:
    line = ""
    for candidate in sample.splitlines():
        if candidate.strip():
            line = candidate
            break
    counts = {"\t": line.count("\t"), ";": line.count(";"), ",": line.count(",")}
    delimiter, count = max(counts.items(), key=lambda item: item[1])
    if count == 0:
        return ","
    return delimiter


def decode_table(data: bytes) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1251")


def read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("В книге нет листов.")
        raw_rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [cell_str(value) for value in row]
            if any(cells):
                raw_rows.append(cells)
    finally:
        workbook.close()
    return table_from_rows(raw_rows)


def read_delimited(text: str) -> tuple[list[str], list[dict[str, str]]]:
    if not text.strip():
        raise ValueError("Файл пустой.")
    delimiter = detect_delimiter(text[:8192])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    return table_from_rows(raw_rows)


def read_table(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if path.suffix.lower() == ".xlsx":
        return read_xlsx(path)
    data = path.read_bytes()
    if data.startswith(b"PK"):
        raise ValueError("Похоже на Excel. Сохраните файл как .xlsx.")
    return read_delimited(decode_table(data))


def score_header(header: str, keywords: tuple[str, ...]) -> int:
    folded = header.casefold()
    best = 0
    for keyword in keywords:
        if folded == keyword:
            best = max(best, 3)
        elif keyword in folded:
            best = max(best, 1)
    return best


def guess_mapping(fields: tuple[MapField, ...], headers: list[str]) -> dict[str, str]:
    used: set[str] = set()
    chosen: dict[str, str] = {}
    for field in fields:
        selected = ""
        best = 0
        for header in headers:
            if header in used:
                continue
            score = score_header(header, field.guesses)
            if score > best:
                best = score
                selected = header
        if selected:
            used.add(selected)
        chosen[field.key] = selected
    return chosen


def parse_mapping(form, fields: tuple[MapField, ...], headers: list[str]) -> dict[str, str]:
    header_set = set(headers)
    mapping: dict[str, str] = {}
    used: set[str] = set()
    for field in fields:
        raw = str(form.get(field.key) or "").strip()
        if not raw:
            if field.required:
                raise ValueError(f"Выберите столбец «{field.label}».")
            mapping[field.key] = ""
            continue
        if raw not in header_set:
            raise ValueError("Такого столбца в файле нет.")
        if raw in used:
            raise ValueError("Каждый столбец можно выбрать только один раз.")
        used.add(raw)
        mapping[field.key] = raw
    return mapping


def mapped_rows(rows: list[dict[str, str]], mapping: dict[str, str]) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    for row in rows:
        item: dict[str, str] = {}
        for key, header in mapping.items():
            item[key] = row.get(header, "").strip() if header else ""
        result.append(item)
    return result
