"""Slow one-tab collector for the logged-in MyHeritage DNA match list."""

from __future__ import annotations

import csv
import os
import random
import re
import subprocess
import sys
import time
from pathlib import Path
from urllib.parse import parse_qs, urlparse

CDP = "http://127.0.0.1:9333"
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "myheritage"
DOWNLOADS = OUT / "downloads"
PROGRESS = OUT / "progress.txt"

KIT_PATH = "OYYV6LTKKSUF6NUSYWBFTVD4YGDX25I"
MEMBER_ID = "OYYV643NWHM2CN2RGIHAGMU4TMJWFVY"
KIT_ID = "9B20E87B-83CA-4B0F-B260-EBA6B5872E96"

PAIR_RE = re.compile(
    r"/dna/match/(D-([0-9A-Fa-f-]{36})-D-([0-9A-Fa-f-]{36}))",
    re.I,
)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"\+\d[\d\s().-]{7,}\d")

MATCH_FIELDS = [
    "match_id",
    "pair_id",
    "name",
    "age_band",
    "country",
    "managed_by",
    "estimated_relationship",
    "shared_percent",
    "total_cm",
    "shared_segments",
    "largest_segment_cm",
    "tree_people",
    "tree_manager",
    "tree_place",
    "tree_status",
    "download_status",
    "list_page",
    "notes",
]
PROB_FIELDS = ["match_id", "relationship", "probability", "mrca"]
SHARED_FIELDS = [
    "match_id",
    "shared_match_id",
    "name",
    "relationship_to_you",
    "percent_to_you",
    "cm_to_you",
    "relationship_to_match",
    "percent_to_match",
    "cm_to_match",
]
PEDIGREE_FIELDS = ["match_id", "person", "lifespan", "relation", "note"]
SEGMENT_FIELDS = [
    "match_id",
    "chromosome",
    "start_location",
    "end_location",
    "start_rsid",
    "end_rsid",
    "centimorgans",
    "snps",
]


class StopCollection(Exception):
    def __init__(self, reason: str) -> None:
        super().__init__(reason)
        self.reason = reason


def log(message: str) -> None:
    print(message, flush=True)


def scrub(value: str) -> str:
    text = EMAIL_RE.sub("", value or "")
    text = PHONE_RE.sub("", text)
    return " ".join(text.replace("\xa0", " ").split())


def list_url(page_no: int) -> str:
    return (
        "https://www.myheritage.com/dna/matches/"
        f"{KIT_PATH}?memberId={MEMBER_ID}&kitId={KIT_ID}"
        f"&p={page_no}&ps=50&sort=total_shared_segments_length_in_cm"
    )


def human_pause(reason: str) -> None:
    seconds = random.uniform(5, 11.25)
    log(f"pause {seconds:.0f}s before {reason}")
    time.sleep(seconds)


def human_scroll(page) -> None:
    for _ in range(random.randint(4, 7)):
        page.mouse.wheel(0, random.randint(180, 420))
        page.wait_for_timeout(random.randint(400, 900))


def read_progress() -> dict[str, str]:
    data = {"page": "1", "matches": "0", "status": "new", "last_match_id": ""}
    if not PROGRESS.exists():
        return data
    for line in PROGRESS.read_text(encoding="utf-8").splitlines():
        if "=" in line:
            key, value = line.split("=", 1)
            data[key.strip()] = value.strip()
    return data


def write_progress(**values: object) -> None:
    current = read_progress()
    current.update({key: str(value) for key, value in values.items()})
    lines = [f"{key}={current[key]}" for key in (
        "status",
        "page",
        "last_match_id",
        "matches",
        "probabilities",
        "shared_matches",
        "pedigrees",
        "segments",
    ) if key in current or key in values]
    # keep any extra keys already stored
    known = {line.split("=", 1)[0] for line in lines}
    for key, value in current.items():
        if key not in known:
            lines.append(f"{key}={value}")
    PROGRESS.write_text("\n".join(lines) + "\n", encoding="utf-8")


def load_ids() -> set[str]:
    path = OUT / "matches.csv"
    if not path.exists() or path.stat().st_size == 0:
        return set()
    with path.open(encoding="utf-8", newline="") as handle:
        return {row["match_id"] for row in csv.DictReader(handle) if row.get("match_id")}


def purge_match(path: Path, match_id: str) -> None:
    if not path.exists() or path.stat().st_size == 0:
        return
    with path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = reader.fieldnames or []
        rows = [row for row in reader if row.get("match_id") != match_id]
    if not fields:
        return
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
        handle.flush()
        os.fsync(handle.fileno())


def append_rows(path: Path, fields: list[str], rows: list[dict]) -> None:
    if not rows:
        return
    new_file = not path.exists() or path.stat().st_size == 0
    with path.open("a", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        if new_file:
            writer.writeheader()
        for row in rows:
            writer.writerow({key: scrub(str(row.get(key, ""))) for key in fields})
        handle.flush()
        os.fsync(handle.fileno())


def count_rows(path: Path) -> int:
    if not path.exists() or path.stat().st_size == 0:
        return 0
    with path.open(encoding="utf-8", newline="") as handle:
        return max(sum(1 for _ in handle) - 1, 0)


def wall(page) -> str:
    try:
        text = page.inner_text("body")
        title = page.title()
    except Exception as exc:
        return f"unreadable:{exc.__class__.__name__}"
    low = f"{title}\n{text}".lower()
    if any(phrase in low for phrase in (
        "security check",
        "i am human",
        "protected by imperva",
        "verify you are human",
    )):
        if "review dna match" not in low and "showing 1" not in low:
            return "security"
    if any(phrase in low for phrase in ("too many requests", "rate limit", "try again later")):
        return "rate"
    try:
        password = page.locator("input[type='password']").count()
    except Exception:
        password = 0
    if password and "review dna match" not in low and "dna matches" not in low:
        return "login"
    return ""


def raise_if_wall(page) -> None:
    reason = wall(page)
    if reason:
        raise StopCollection(reason)


def wait_ready(page) -> None:
    # Load state and locator waits only. wait_for_function and string eval
    # hit EvalError because MyHeritage CSP forbids unsafe-eval.
    try:
        page.wait_for_load_state("domcontentloaded", timeout=60000)
    except Exception as exc:
        raise StopCollection("page not ready") from exc
    ready = page.locator(
        "a[href*='/dna/match/'], "
        "[data-automations='MatchProfileDetailsOther'], "
        ".pedigree_tree_person, "
        ".pedigree_tree, "
        "input[type='password']"
    ).or_(page.get_by_text(
        re.compile(
            r"security check|i am human|protected by imperva|verify you are human|"
            r"pedigree|sign in|too many requests|request membership|family tree|dna match",
            re.I,
        )
    ))
    try:
        ready.first.wait_for(state="attached", timeout=50000)
    except Exception as exc:
        raise StopCollection("page not ready") from exc


def goto(page, url: str) -> None:
    last: Exception | None = None
    for attempt in range(3):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=60000)
            wait_ready(page)
            raise_if_wall(page)
            return
        except StopCollection:
            raise
        except Exception as exc:
            last = exc
            log(f"goto retry {attempt + 1}: {type(exc).__name__}")
            human_pause("retry")
    raise StopCollection("page not ready") from last


def match_links(page) -> list[tuple[str, str, str]]:
    hrefs = page.evaluate(
        """() => [...document.querySelectorAll('a[href*="/dna/match/"]')].map(a => a.href)"""
    )
    found: list[tuple[str, str, str]] = []
    seen: set[str] = set()
    for href in hrefs:
        match = PAIR_RE.search(href or "")
        if not match:
            continue
        pair_id, _kit, match_id = match.group(1), match.group(2), match.group(3).upper()
        if match_id in seen:
            continue
        seen.add(match_id)
        found.append((match_id, pair_id, href.split("#")[0]))
    return found


def first_number(text: str, pattern: str) -> str:
    match = re.search(pattern, text or "", re.I)
    return match.group(1).replace(",", ".") if match else ""


def profile_fields(page) -> dict:
    raw = page.evaluate(
        """() => {
          const root = document.querySelector('[data-automations="MatchProfileDetailsOther"]');
          const textOf = (sel, base=document) => {
            const el = base.querySelector(sel);
            return el ? el.innerText.replace(/\\s+/g, ' ').trim() : '';
          };
          const details = root ? [...root.querySelectorAll('.detail_property')].map(el => el.innerText.replace(/\\s+/g, ' ').trim()) : [];
          return {
            name: root ? textOf('[data-automations="ProfileDetailsName"]', root) : '',
            details,
            tree: root ? textOf('.matched_family_tree_details', root) : '',
            treeHref: root ? (root.querySelector('a.view_tree')?.href || '') : '',
            relationship: textOf('[data-automations="possible_relationships_static_text"]'),
            shared: textOf('[data-automations="QualitySharedDnaValue"]'),
            segments: textOf('[data-automations="QualitySharedSegmentsTotal"]'),
            largest: textOf('[data-automations="QualityLargestSegmentTotal"]'),
            obfuscated: /"other_dna_kit"\\s*:\\s*\\{\\s*"is_pedigree_chart_obfuscated"\\s*:\\s*true/.test(document.documentElement.innerHTML),
            hiddenPhrase: (document.body.innerText || '').toLowerCase().includes('may reveal new information')
          };
        }"""
    )
    age = country = managed = ""
    for detail in raw["details"]:
        low = detail.lower()
        if low.startswith("age"):
            age = detail.split(":", 1)[-1].strip()
        elif low.startswith("from"):
            country = detail.split(":", 1)[-1].strip()
        elif "managed by" in low:
            managed = re.sub(r"(?i).*managed by", "", detail).strip()
    tree = raw["tree"]
    people = first_number(tree, r"with\s+(\d+)\s+people")
    place = ""
    place_match = re.search(r"\bfrom\s+(.+)$", tree, re.I)
    if place_match:
        place = place_match.group(1).strip(" .")
    manager = ""
    manager_match = re.search(r"managed by\s+(.+?)\s+from\b", tree, re.I)
    if manager_match:
        manager = manager_match.group(1).strip()
    return {
        "name": raw["name"],
        "age_band": age,
        "country": country,
        "managed_by": managed,
        "estimated_relationship": raw["relationship"],
        "shared_percent": first_number(raw["shared"], r"([\d.,]+)\s*%"),
        "total_cm": first_number(raw["shared"], r"([\d.,]+)\s*(?:‎\s*)?cM"),
        "shared_segments": first_number(raw["segments"], r"(\d+)"),
        "largest_segment_cm": first_number(raw["largest"], r"([\d.,]+)\s*(?:‎\s*)?cM"),
        "tree_people": people,
        "tree_manager": manager,
        "tree_place": place,
        "tree_href": raw["treeHref"],
        "tree_hidden": bool(raw["obfuscated"] or raw["hiddenPhrase"]),
    }


def probability_rows(page, match_id: str) -> list[dict]:
    raw = page.evaluate(
        """() => [...document.querySelectorAll('[data-automations="probable_relationship_cell_relationship_cell"]')].map(el => {
          const row = el.parentElement;
          const text = (node) => node ? node.innerText.replace(/\\s+/g, ' ').trim() : '';
          return {
            relationship: text(el),
            probability: text(row && row.querySelector('[data-automations="probable_relationship_cell_probability_cell"]')),
            mrca: text(row && row.querySelector('[data-automations="probable_relationship_cell_mrca_cell"]'))
          };
        })"""
    )
    return [{"match_id": match_id, **row} for row in raw if row.get("relationship")]


def shared_rows(page, match_id: str) -> list[dict]:
    raw = page.evaluate(
        """() => [...document.querySelectorAll('.dna_shared_matches_row')].map(row => {
          const text = (sel) => {
            const el = row.querySelector(`[data-automations="${sel}"]`);
            return el ? el.innerText.replace(/\\s+/g, ' ').trim() : '';
          };
          const box = row.querySelector('[data-automations^="match_checkbox_dnamatch-"]');
          return {
            name: text('SharedMatchProfileName'),
            rel_you: text('SharedMatchPossibleRelationships'),
            pct_you: text('SharedMatchPer'),
            cm_you: text('SharedMatchTotal'),
            rel_other: text('SharedMatchOtherPossibleRelationships'),
            pct_other: text('SharedMatchOtherPer'),
            cm_other: text('SharedMatchOtherTotal'),
            key: box ? box.getAttribute('data-automations') : ''
          };
        })"""
    )
    rows = []
    for item in raw:
        if not item.get("name"):
            continue
        found = PAIR_RE.search(item.get("key") or "")
        rows.append({
            "match_id": match_id,
            "shared_match_id": found.group(3).upper() if found else "",
            "name": item["name"],
            "relationship_to_you": item["rel_you"],
            "percent_to_you": item["pct_you"],
            "cm_to_you": item["cm_you"],
            "relationship_to_match": item["rel_other"],
            "percent_to_match": item["pct_other"],
            "cm_to_match": item["cm_other"],
        })
    return rows


def click_visible(page, selector: str) -> bool:
    loc = page.locator(selector)
    try:
        count = min(loc.count(), 6)
    except Exception:
        return False
    for index in range(count):
        item = loc.nth(index)
        try:
            if item.is_visible():
                item.scroll_into_view_if_needed(timeout=5000)
                item.click(timeout=8000)
                return True
        except Exception:
            continue
    return False


def expand_relationships(page) -> None:
    button = page.locator("button[data-automations='collapsible_toggle_button']", has_text="Show more relationships")
    try:
        if button.count() and button.first.is_visible():
            human_pause("show-more")
            button.first.scroll_into_view_if_needed(timeout=5000)
            button.first.click(timeout=8000)
            page.wait_for_timeout(1500)
    except Exception as exc:
        log(f"relationships show-more skipped: {exc.__class__.__name__}")


def expand_shared(page) -> str:
    notes = ""
    for _ in range(20):
        loc = page.locator("a[data-automations='SharedMatchesShowMore']")
        try:
            if not loc.count() or not loc.first.is_visible():
                break
        except Exception:
            break
        before = page.locator(".dna_shared_matches_row").count()
        human_pause("show-more")
        raise_if_wall(page)
        if not click_visible(page, "a[data-automations='SharedMatchesShowMore']"):
            break
        grew = False
        for _wait in range(12):
            page.wait_for_timeout(1000)
            if page.locator(".dna_shared_matches_row").count() > before:
                grew = True
                break
        if not grew:
            notes = "show-more stalled"
            break
    return notes


def download_shared(page, match_id: str) -> str:
    unlock = page.locator("[data-automations='action-obfuscation-mask-button-SharedSegmentsUnlock']")
    try:
        locked = unlock.count() and unlock.first.is_visible()
    except Exception:
        locked = False
    human_pause("download")
    raise_if_wall(page)
    if not click_visible(page, "button[aria-label*='Advanced options']"):
        return "locked" if locked else "absent"
    item = page.get_by_text("Download shared DNA info", exact=True)
    try:
        item.first.wait_for(state="visible", timeout=8000)
    except Exception:
        page.keyboard.press("Escape")
        return "locked" if locked else "absent"
    href = ""
    try:
        href = item.first.evaluate("el => el.href || (el.closest('a') && el.closest('a').href) || ''")
    except Exception:
        href = ""
    if "paywall" in href or "subscription" in href:
        page.keyboard.press("Escape")
        return "paywall"
    dest_base = DOWNLOADS / match_id
    try:
        with page.expect_download(timeout=30000) as caught:
            item.first.click(timeout=8000)
        download = caught.value
        suffix = Path(download.suggested_filename or "").suffix or ".csv"
        target = dest_base.with_suffix(suffix.lower())
        download.save_as(str(target))
    except Exception as exc:
        log(f"download failed: {exc.__class__.__name__}")
        if "paywall" in page.url or "subscription" in page.url:
            return "paywall"
        return "failed"
    data = target.read_bytes()[:200].lstrip().lower()
    if data.startswith(b"<!doctype") or data.startswith(b"<html"):
        target.unlink(missing_ok=True)
        return "paywall"
    return "saved"


def header_key(name: str) -> str:
    lowered = re.sub(r"[^a-z]", "", (name or "").lower())
    if "chrom" in lowered:
        return "chromosome"
    if lowered in {"startlocation", "start", "begin"}:
        return "start_location"
    if lowered in {"endlocation", "end", "stop"}:
        return "end_location"
    if "startrsid" in lowered or lowered == "startrs":
        return "start_rsid"
    if "endrsid" in lowered or lowered == "endrs":
        return "end_rsid"
    if "centimorgan" in lowered or lowered in {"cm", "cmvalue"}:
        return "centimorgans"
    if "snp" in lowered:
        return "snps"
    return ""


def rows_from_table(headers: list[str], records: list[list[str]], match_id: str) -> list[dict]:
    keys = [header_key(header) for header in headers]
    if "chromosome" not in keys:
        return []
    parsed = []
    for record in records:
        row = {"match_id": match_id}
        for key, value in zip(keys, record):
            if key:
                row[key] = value.strip()
        if row.get("chromosome"):
            parsed.append(row)
    return parsed


def parse_segments(path: Path, match_id: str) -> list[dict]:
    raw = path.read_bytes()
    if raw[:2] == b"PK":
        try:
            from openpyxl import load_workbook
        except Exception:
            log("xlsx download left unparsed")
            return []
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book.active
        table = [[("" if cell is None else str(cell)) for cell in row] for row in sheet.iter_rows(values_only=True)]
        book.close()
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:2000]
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
        table = list(csv.reader(text.splitlines(), delimiter=delimiter))
    header_at = None
    for index, record in enumerate(table[:25]):
        if any(header_key(cell) == "chromosome" for cell in record):
            header_at = index
            break
    if header_at is None:
        return []
    return rows_from_table(table[header_at], table[header_at + 1:], match_id)


def _pedigree_blank(note: str) -> list[dict]:
    return [{"person": "", "lifespan": "", "relation": "", "note": note}]


def _pedigree_people(page) -> list[str]:
    loc = page.locator(".pedigree_tree_person")
    people = []
    for index in range(loc.count()):
        person = " ".join((loc.nth(index).inner_text(timeout=3000) or "").split())
        if person:
            people.append(person)
    return people


def _open_pedigree(page, tree_href: str) -> tuple[str, list[dict]]:
    goto(page, tree_href)
    human_scroll(page)
    text = page.locator("body").inner_text(timeout=15000)
    low = text.lower()
    restricted = any(phrase in low for phrase in (
        "you are not a member",
        "may be restricted",
        "this tree is private",
        "tree is private",
        "request membership",
    ))
    people = _pedigree_people(page)
    named = [person for person in people if person.casefold() not in {"unknown", "add", "add person"}]
    count = re.search(r"(\d+)\s+of\s+(\d+)\s+people", text, re.I)
    if restricted and not named:
        note = "tree restricted"
        if count:
            note = f"tree restricted, {count.group(1)} of {count.group(2)} shown"
        return "hidden", _pedigree_blank(note)
    rows = [{"person": person, "lifespan": "", "relation": "", "note": "pedigree chart"} for person in named]
    if rows:
        return "visible", rows
    if people:
        note = "pedigree chart showed no named people"
        if count:
            note = f"{note}, {count.group(1)} of {count.group(2)}"
        return "hidden", _pedigree_blank(note)
    return "absent", []


def read_pedigree(page, tree_href: str, hidden: bool) -> tuple[str, list[dict]]:
    if hidden:
        return "hidden", _pedigree_blank("tree hidden on the match page")
    if not tree_href:
        return "absent", []
    human_pause("pedigree")
    try:
        return _open_pedigree(page, tree_href)
    except StopCollection as exc:
        if exc.reason in {"security", "rate", "login"}:
            raise
        log(f"pedigree skipped: {exc.reason}")
        return "error", _pedigree_blank("pedigree unavailable")
    except Exception as exc:
        log(f"pedigree skipped: {exc.__class__.__name__}")
        return "error", _pedigree_blank("pedigree unavailable")


def save_match(match_id: str, match_row: dict, probs: list[dict], shared: list[dict], pedigree: list[dict], segments: list[dict]) -> None:
    for path in (
        OUT / "matches.csv",
        OUT / "probabilities.csv",
        OUT / "shared_matches.csv",
        OUT / "pedigrees.csv",
        OUT / "segments.csv",
    ):
        purge_match(path, match_id)
    append_rows(OUT / "probabilities.csv", PROB_FIELDS, probs)
    append_rows(OUT / "shared_matches.csv", SHARED_FIELDS, shared)
    append_rows(OUT / "pedigrees.csv", PEDIGREE_FIELDS, pedigree)
    append_rows(OUT / "segments.csv", SEGMENT_FIELDS, segments)
    append_rows(OUT / "matches.csv", MATCH_FIELDS, [match_row])


def collect_one(page, match_id: str, pair_id: str, url: str, list_page: int) -> None:
    human_pause("match page")
    goto(page, url)
    human_scroll(page)
    raise_if_wall(page)
    try:
        page.wait_for_selector("[data-automations='MatchProfileDetailsOther']", timeout=20000)
    except Exception:
        raise_if_wall(page)
        raise StopCollection("match page had no profile")
    fields = profile_fields(page)
    if not fields["name"]:
        raise StopCollection("match page had no name")
    expand_relationships(page)
    note = expand_shared(page)
    probs = probability_rows(page, match_id)
    shared = shared_rows(page, match_id)
    download_status = download_shared(page, match_id)
    segments: list[dict] = []
    saved_file = next(DOWNLOADS.glob(f"{match_id}.*"), None)
    if download_status == "saved" and saved_file:
        segments = parse_segments(saved_file, match_id)
        if not segments:
            download_status = "saved-unparsed"
    tree_status, pedigree = read_pedigree(page, fields["tree_href"], fields["tree_hidden"])
    pedigree_rows = [{"match_id": match_id, **row} for row in pedigree]
    notes = " ".join(part for part in (note, f"segments:{len(segments)}" if segments else "") if part)
    match_row = {
        "match_id": match_id,
        "pair_id": pair_id,
        "name": fields["name"],
        "age_band": fields["age_band"],
        "country": fields["country"],
        "managed_by": fields["managed_by"],
        "estimated_relationship": fields["estimated_relationship"],
        "shared_percent": fields["shared_percent"],
        "total_cm": fields["total_cm"],
        "shared_segments": fields["shared_segments"],
        "largest_segment_cm": fields["largest_segment_cm"],
        "tree_people": fields["tree_people"],
        "tree_manager": fields["tree_manager"],
        "tree_place": fields["tree_place"],
        "tree_status": tree_status,
        "download_status": download_status,
        "list_page": str(list_page),
        "notes": notes,
    }
    save_match(match_id, match_row, probs, shared, pedigree_rows, segments)
    write_progress(
        status="running",
        page=list_page,
        last_match_id=match_id,
        matches=count_rows(OUT / "matches.csv"),
        probabilities=count_rows(OUT / "probabilities.csv"),
        shared_matches=count_rows(OUT / "shared_matches.csv"),
        pedigrees=count_rows(OUT / "pedigrees.csv"),
        segments=count_rows(OUT / "segments.csv"),
    )
    log(f"saved match {match_id} download={download_status} tree={tree_status} shared={len(shared)}")


def on_list(url: str, page_no: int) -> bool:
    parsed = urlparse(url or "")
    if "/dna/matches/" not in parsed.path:
        return False
    current = parse_qs(parsed.query).get("p", ["1"])[0]
    return current == str(page_no)


def other_collector_running() -> bool:
    """True only for a real python process running this script, not a shell that mentions it."""
    me = {os.getpid(), os.getppid()}
    try:
        output = subprocess.check_output(["ps", "-ax", "-o", "pid=,command="], text=True)
    except subprocess.CalledProcessError:
        return False
    for line in output.splitlines():
        line = line.strip()
        if not line or "myheritage_collect.py" not in line:
            continue
        pid_s, _, cmd = line.partition(" ")
        try:
            pid = int(pid_s)
        except ValueError:
            continue
        if pid in me or "<defunct>" in cmd:
            continue
        # Skip shells / wrappers whose command line merely quotes the script path.
        if any(marker in cmd for marker in ("/bin/zsh", "/bin/bash", "/bin/sh", " -c ", "builtin eval")):
            continue
        tokens = cmd.split()
        if not tokens:
            continue
        interpreter = Path(tokens[0]).name
        if not interpreter.startswith("python"):
            continue
        if any(token.endswith("myheritage_collect.py") for token in tokens):
            return True
    return False


def myheritage_page(browser):
    found = []
    for context in browser.contexts:
        for page in context.pages:
            if "myheritage.com" in (page.url or ""):
                found.append(page)
    if not found:
        raise StopCollection("no myheritage tab")
    return found[0]


def run() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    if other_collector_running():
        log("collector already running")
        return 0
    from playwright.sync_api import sync_playwright

    progress = read_progress()
    page_no = int(progress.get("page") or "1")
    done = load_ids()
    write_progress(status="running", page=page_no, matches=len(done))
    log(f"start page={page_no} already={len(done)}")
    playwright = sync_playwright().start()
    try:
        browser = playwright.chromium.connect_over_cdp(CDP)
        page = myheritage_page(browser)
        page.set_default_timeout(30000)

        def close_extra(new_page) -> None:
            if new_page is not page:
                log("closed extra tab")
                try:
                    new_page.close()
                except Exception:
                    pass

        for context in browser.contexts:
            context.on("page", close_extra)

        while True:
            url = list_url(page_no)
            if not on_list(page.url, page_no):
                human_pause("pagination")
                goto(page, url)
            else:
                raise_if_wall(page)
            human_scroll(page)
            links = match_links(page)
            fresh = [item for item in links if item[0] not in done]
            log(f"list page {page_no}: links={len(links)} new={len(fresh)}")
            if not fresh:
                write_progress(status="done", page=page_no, matches=len(done))
                log("no new matches on this page")
                return 0
            list_target = url
            for match_id, pair_id, match_url in fresh:
                try:
                    collect_one(page, match_id, pair_id, match_url, page_no)
                    done.add(match_id)
                except StopCollection:
                    raise
                except Exception as exc:
                    log(f"skip {match_id}: {type(exc).__name__}")
                human_pause("list")
                goto(page, list_target)
                human_scroll(page)
            page_no += 1
            write_progress(status="running", page=page_no, matches=len(done))
    except StopCollection as exc:
        write_progress(status=f"stopped:{exc.reason}", matches=len(load_ids()))
        log(f"STOP {exc.reason}")
        return 3
    finally:
        playwright.stop()


if __name__ == "__main__":
    try:
        sys.exit(run())
    except KeyboardInterrupt:
        write_progress(status="paused", matches=len(load_ids()))
        log("STOP paused")
        sys.exit(0)
